import datetime
import difflib
import glob
import io
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from functools import lru_cache

import pymupdf

from . import dni
from .i18n import N_, _

BASE14 = {"Helvetica": "helv", "Times": "tiro", "Courier": "cour"}
IMAGE_EXT = (".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp", ".pnm", ".svg")
OFFICE_EXT = (".doc", ".docx", ".odt", ".rtf", ".txt", ".xls", ".xlsx", ".ods", ".csv",
              ".ppt", ".pptx", ".odp", ".html", ".htm", ".odg")
EBOOK_EXT = (".xps", ".oxps", ".epub", ".cbz", ".fb2", ".mobi")


class OperationError(Exception):
    pass


@lru_cache(maxsize=256)
def font_file(family, bold=False, italic=False):
    style = " ".join(s for s, on in (("Bold", bold), ("Italic", italic)) if on)
    pattern = family + (f":style={style}" if style else "")
    try:
        out = subprocess.run(["fc-match", "-f", "%{file}", pattern],
                             capture_output=True, text=True, timeout=5).stdout.strip()
    except (OSError, subprocess.SubprocessError):
        return None
    if out.lower().endswith((".ttf", ".otf", ".ttc")) and os.path.exists(out):
        return out
    return None


def font_args(family, bold=False, italic=False):
    path = font_file(family, bold, italic)
    if path:
        tag = "F" + "".join(c for c in os.path.basename(path) if c.isalnum())[:24]
        return {"fontname": tag, "fontfile": path}
    return {"fontname": "helv"}


def font_object(family, bold=False, italic=False):
    path = font_file(family, bold, italic)
    return pymupdf.Font(fontfile=path) if path else pymupdf.Font("helv")


def to_unrotated(page, rect):
    return pymupdf.Rect(rect) * page.derotation_matrix


def to_unrotated_pt(page, pt):
    return pymupdf.Point(pt) * page.derotation_matrix


def to_visible(page, rect):
    return pymupdf.Rect(rect) * page.rotation_matrix


def is_svg(data):
    head = data[:2048].lstrip().lower()
    return head.startswith(b"<svg") or (head.startswith(b"<?xml") and b"<svg" in head) or b"<svg" in head[:512]


def svg_document(data):
    src = pymupdf.open(stream=data, filetype="svg")
    pdf = pymupdf.open("pdf", src.convert_to_pdf())
    src.close()
    return pdf


def graphic_size(data):
    if is_svg(data):
        doc = svg_document(data)
        r = doc[0].rect
        doc.close()
        return r.width, r.height
    pix = pymupdf.Pixmap(data)
    return pix.width, pix.height


_SVG_COLOR = re.compile(r'((?:fill|stroke)\s*[:=]\s*["\']?)(#[0-9a-fA-F]{3,8}|rgb\([^)]*\)|[a-zA-Z]+)')


def recolor_svg(data, color_hex):
    text = data.decode("utf-8", "replace")

    def repl(m):
        if m.group(2).lower() in ("none", "transparent", "currentcolor", "inherit", "url"):
            return m.group(0)
        return m.group(1) + color_hex
    text = _SVG_COLOR.sub(repl, text)
    text = re.sub(r"<svg\b", f'<svg color="{color_hex}"', text, count=1)
    return text.encode("utf-8")


def recolor_png(data, rgb):
    pix = pymupdf.Pixmap(data)
    if pix.colorspace is None or pix.colorspace.n != 3:
        pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
    if not pix.alpha:
        pix = pymupdf.Pixmap(pix, 1)
    r, g, b = (int(c * 255) for c in rgb)
    samples = bytearray(pix.samples)
    n = pix.n
    for i in range(0, len(samples), n):
        lum = (samples[i] * 299 + samples[i + 1] * 587 + samples[i + 2] * 114) // 1000
        if samples[i + n - 1] and lum < 235:
            samples[i], samples[i + 1], samples[i + 2] = r, g, b
        else:
            samples[i + n - 1] = 0
    out = pymupdf.Pixmap(pix.colorspace, pix.width, pix.height, bytes(samples), True)
    return out.tobytes("png")


def optimize_image(data, width_pt, height_pt, dpi=300, jpeg_quality=85):
    from PIL import Image
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception:
        return data
    max_w = max(1, round(width_pt / 72 * dpi))
    max_h = max(1, round(height_pt / 72 * dpi))
    if img.width > max_w or img.height > max_h:
        img.thumbnail((max_w, max_h), Image.Resampling.LANCZOS)
    has_alpha = img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info)
    out = io.BytesIO()
    if has_alpha:
        img.convert("RGBA" if img.mode != "LA" else "LA").save(out, "PNG", optimize=True)
    else:
        img.convert("L" if img.mode in ("L", "1") else "RGB").save(out, "JPEG", quality=jpeg_quality,
                                                                  optimize=True)
    result = out.getvalue()
    return result if len(result) < len(data) else data


GRAPHIC_RECT = "ACRPDFRect"
GRAPHIC_SOURCE = "ACRPDFSource"


def _unique_png(data):
    import struct
    import uuid
    import zlib
    if not data.startswith(b"\x89PNG") or not data.endswith(b"IEND\xaeB`\x82"):
        return data
    payload = b"acrpdf\x00" + uuid.uuid4().hex.encode()
    chunk = struct.pack(">I", len(payload)) + b"tEXt" + payload + struct.pack(">I", zlib.crc32(b"tEXt" + payload))
    return data[:-12] + chunk + data[-12:]


def insert_graphic(page, rect, data, track=True):
    target = to_unrotated(page, rect)
    if is_svg(data):
        src = svg_document(data)
        xref = page.show_pdf_page(target, src, 0, keep_proportion=True, rotate=page.rotation, overlay=True)
        src.close()
    else:
        image = optimize_image(data, abs(target.width), abs(target.height))
        image = _unique_png(image)
        xref = page.insert_image(target, stream=image, keep_proportion=True, rotate=page.rotation, overlay=True)
    if track and xref:
        doc = page.parent
        source = doc.get_new_xref()
        doc.update_object(source, "<<>>")
        doc.update_stream(source, data)
        r = pymupdf.Rect(rect)
        doc.xref_set_key(xref, GRAPHIC_RECT, f"[{r.x0:.3f} {r.y0:.3f} {r.x1:.3f} {r.y1:.3f}]")
        doc.xref_set_key(xref, GRAPHIC_SOURCE, f"{source} 0 R")
    return xref


def placed_graphics(page):
    doc = page.parent
    candidates = [img[0] for img in page.get_images(full=True)] + [x[0] for x in page.get_xobjects()]
    found = []
    for xref in dict.fromkeys(candidates):
        kind, value = doc.xref_get_key(xref, GRAPHIC_RECT)
        if kind == "array":
            found.append((xref, pymupdf.Rect([float(v) for v in value.strip("[]").split()])))
    return found


def graphic_at(page, point):
    pt = pymupdf.Point(point)
    hits = [(xref, r) for xref, r in placed_graphics(page) if r.contains(pt)]
    return hits[-1] if hits else None


def graphic_source(doc, xref):
    kind, value = doc.xref_get_key(xref, GRAPHIC_SOURCE)
    return doc.xref_stream(int(value.split()[0])) if kind == "xref" else None


def remove_graphic(page, xref):
    doc = page.parent
    if doc.xref_get_key(xref, "Subtype")[1] == "/Image":
        page.delete_image(xref)
    else:
        doc.update_stream(xref, b" ")
    doc.xref_set_key(xref, GRAPHIC_RECT, "null")


def move_graphic(page, xref, rect):
    data = graphic_source(page.parent, xref)
    if data is None:
        return None
    remove_graphic(page, xref)
    return insert_graphic(page, rect, data)


def render_graphic_png(data, width_px=600):
    if is_svg(data):
        doc = svg_document(data)
        zoom = width_px / max(doc[0].rect.width, 1)
        png = doc[0].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=True).tobytes("png")
        doc.close()
        return png
    return data


def insert_text_block(page, rect, text, family="Helvetica", size=12, color=(0, 0, 0),
                      bold=False, italic=False, align=0):
    kw = font_args(family, bold, italic)
    font = font_object(family, bold, italic)
    lines = text.split("\n")
    width = max(font.text_length(line, fontsize=size) for line in lines) + 2
    height = size * 1.25 * len(lines) + size * 0.4
    r = pymupdf.Rect(rect.x0, rect.y0, max(rect.x1, rect.x0 + width), max(rect.y1, rect.y0 + height))
    box = to_unrotated(page, r)
    rc = page.insert_textbox(box, text, fontsize=size, color=color, align=align,
                             rotate=page.rotation, **kw)
    if rc < 0:
        box = to_unrotated(page, pymupdf.Rect(r.x0, r.y0, r.x0 + width + 50, r.y0 + height * 2 + 20))
        page.insert_textbox(box, text, fontsize=size, color=color, align=align,
                            rotate=page.rotation, **kw)


def find_span(page, point):
    pt = pymupdf.Point(point)
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                if to_visible(page, span["bbox"]).contains(pt) and span["text"].strip():
                    return span
    return None


def is_ocr_span(span):
    return span.get("alpha", 255) == 0 or span.get("font") == "GlyphLessFont"


def _background_color(page, rect):
    pix = page.get_pixmap(clip=pymupdf.Rect(rect) + (-4, -4, 4, 4), dpi=72, alpha=False)
    samples = pix.samples
    n = pix.n
    pixels = [tuple(samples[i:i + 3]) for i in range(0, len(samples), n)]
    if not pixels:
        return (1, 1, 1)
    light = sorted(pixels, key=sum)[len(pixels) * 3 // 4:]
    return tuple(sum(c[k] for c in light) / len(light) / 255 for k in range(3))


def replace_span(page, span, new_text, family=None, size=None, color=None):
    bbox = pymupdf.Rect(span["bbox"])
    ocr = is_ocr_span(span)
    if ocr:
        page.add_redact_annot(bbox + (-1, -1, 1, 1), fill=_background_color(page, to_visible(page, bbox)))
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
    else:
        page.add_redact_annot(bbox, fill=None)
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
    bold, italic = bool(span["flags"] & 16), bool(span["flags"] & 2)
    if color is None:
        c = span["color"]
        color = (0.1, 0.1, 0.1) if ocr else ((c >> 16 & 255) / 255, (c >> 8 & 255) / 255, (c & 255) / 255)
    if not family:
        family = "DejaVu Serif" if ocr else span["font"].split("+")[-1].split("-")[0]
    page.insert_text(pymupdf.Point(span["origin"]), new_text, fontsize=size or span["size"],
                     color=color, rotate=page.rotation, **font_args(family, bold, italic))


def words_in_rect(page, rect):
    r = to_unrotated(page, rect)
    return [w for w in page.get_text("words") if pymupdf.Rect(w[:4]).intersects(r)]


def markup(page, rect, kind, color):
    words = words_in_rect(page, rect)
    if not words:
        return None
    quads = [pymupdf.Rect(w[:4]).quad for w in words]
    add = {"highlight": page.add_highlight_annot, "underline": page.add_underline_annot,
           "strike": page.add_strikeout_annot, "squiggly": page.add_squiggly_annot}[kind]
    annot = add(quads)
    annot.set_colors(stroke=color)
    annot.update()
    return annot


def shape_annot(page, kind, rect=None, p1=None, p2=None, color=(1, 0, 0), width=2, fill=None):
    if kind == "rect":
        annot = page.add_rect_annot(to_unrotated(page, rect))
    elif kind == "ellipse":
        annot = page.add_circle_annot(to_unrotated(page, rect))
    else:
        annot = page.add_line_annot(to_unrotated_pt(page, p1), to_unrotated_pt(page, p2))
        if kind == "arrow":
            annot.set_line_ends(pymupdf.PDF_ANNOT_LE_NONE, pymupdf.PDF_ANNOT_LE_OPEN_ARROW)
    annot.set_colors(stroke=color, fill=fill)
    annot.set_border(width=width)
    annot.update()
    return annot


def ink_annot(page, strokes, color=(0, 0, 1), width=2):
    annot = page.add_ink_annot([[tuple(to_unrotated_pt(page, p)) for p in s] for s in strokes])
    annot.set_colors(stroke=color)
    annot.set_border(width=width)
    annot.update()
    return annot


def note_annot(page, point, text, author=""):
    annot = page.add_text_annot(to_unrotated_pt(page, point), text, icon="Comment")
    annot.set_info(title=author or os.environ.get("USER", ""), content=text)
    annot.update()
    return annot


def freetext_annot(page, rect, text, size=12, color=(0, 0, 0), family="Helvetica"):
    annot = page.add_freetext_annot(to_unrotated(page, rect), text, fontsize=size,
                                    fontname=BASE14.get(family, "helv"), text_color=color,
                                    rotate=page.rotation)
    annot.update()
    return annot


def annot_at(page, point):
    pt = to_unrotated_pt(page, point)
    for annot in page.annots():
        if annot.rect.contains(pt):
            return annot
    return None


def widget_at(page, point):
    pt = to_unrotated_pt(page, point)
    for w in page.widgets():
        if w.rect.contains(pt):
            return w
    return None


STAMP_PRESETS = [
    (N_("APPROVED"), (0.10, 0.55, 0.20)), (N_("REJECTED"), (0.80, 0.10, 0.10)),
    (N_("CONFIDENTIAL"), (0.80, 0.10, 0.10)), (N_("DRAFT"), (0.20, 0.35, 0.75)),
    (N_("RECEIVED"), (0.20, 0.35, 0.75)), (N_("REVIEWED"), (0.10, 0.55, 0.20)),
    (N_("PAID"), (0.10, 0.55, 0.20)), (N_("COPY"), (0.45, 0.45, 0.45)),
    (N_("FINAL"), (0.10, 0.55, 0.20)), (N_("VOID"), (0.80, 0.10, 0.10)),
]


def draw_stamp(page, rect, title, color=(0.1, 0.55, 0.2), family="DejaVu Sans", subtitle=None,
               opacity=0.9):
    rect = to_unrotated(page, pymupdf.Rect(rect))
    shape = page.new_shape()
    shape.draw_rect(rect, radius=0.12)
    shape.finish(color=color, width=max(1.5, rect.height * 0.05), stroke_opacity=opacity)
    m = rect.height * 0.08
    shape.draw_rect(rect + (m, m, -m, -m), radius=0.1)
    shape.finish(color=color, width=max(0.6, rect.height * 0.015), stroke_opacity=opacity)
    shape.commit()

    max_w = rect.width * 0.86
    size = rect.height * (0.42 if subtitle else 0.5)
    font = font_object(family, bold=True)
    tw = font.text_length(title, fontsize=size)
    if tw > max_w:
        size, tw = size * max_w / tw, max_w
    baseline = rect.y0 + rect.height * (0.44 if subtitle else 0.5) + size * 0.35
    page.insert_text((rect.x0 + (rect.width - tw) / 2, baseline), title, fontsize=size,
                     color=color, fill_opacity=opacity, **font_args(family, bold=True))
    if subtitle:
        sub_size = rect.height * 0.14
        sub_font = font_object(family)
        sw = sub_font.text_length(subtitle, fontsize=sub_size)
        if sw > max_w:
            sub_size, sw = sub_size * max_w / sw, max_w
        page.insert_text((rect.x0 + (rect.width - sw) / 2, rect.y1 - rect.height * 0.2), subtitle,
                         fontsize=sub_size, color=color, fill_opacity=opacity, **font_args(family))


def stamp_png(title, color, family, subtitle, width=190, height=62, zoom=3):
    doc = pymupdf.open()
    page = doc.new_page(width=width, height=height)
    draw_stamp(page, page.rect + (1, 1, -1, -1), title, color, family, subtitle)
    png = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=True).tobytes("png")
    doc.close()
    return png


WATERMARK_POSITIONS = ("center", "top", "bottom", "tile")


def watermark(doc, text=None, image=None, size=60, color=(0.6, 0.6, 0.6), opacity=0.3,
              family="DejaVu Sans", angle=45, pages=None, position="center"):
    font = font_object(family, bold=True)
    kw = font_args(family, bold=True)
    for i in pages or range(doc.page_count):
        page = doc[i]
        r = page.rect
        if position == "tile":
            step_x, step_y = r.width / 2.2, r.height / 3.2
            centers = [pymupdf.Point(step_x * (0.55 + c), step_y * (0.5 + row))
                       for row in range(4) for c in range(3)]
            scale = 0.45
        else:
            y = {"center": r.height / 2, "top": r.height * 0.15, "bottom": r.height * 0.85}[position]
            centers, scale = [pymupdf.Point(r.width / 2, y)], 1.0
        for center in centers:
            if image:
                w = r.width * 0.5 * scale
                iw, ih = graphic_size(image)
                box = pymupdf.Rect(center.x - w / 2, center.y - w * ih / iw / 2,
                                   center.x + w / 2, center.y + w * ih / iw / 2)
                insert_graphic(page, box, image)
                continue
            fs = size * scale
            tw = font.text_length(text, fontsize=fs)
            start = pymupdf.Point(center.x - tw / 2, center.y + fs / 3)
            page.insert_text(start * page.derotation_matrix, text, fontsize=fs, color=color,
                             fill_opacity=opacity,
                             morph=(center * page.derotation_matrix, pymupdf.Matrix(-angle - page.rotation)),
                             overlay=True, **kw)


PAGE_NUMBER_POSITIONS = ("bottom-center", "bottom-right", "bottom-left",
                         "top-center", "top-right", "top-left")


def page_numbers(doc, fmt="{n} / {total}", position="bottom-center", start=1, size=10,
                 family="DejaVu Sans", color=(0, 0, 0), margin=24, first_page=1):
    total = doc.page_count
    today = datetime.date.today().strftime("%d/%m/%Y")
    vpos, hpos = position.split("-")
    align = {"left": 0, "center": 1, "right": 2}[hpos]
    kw = font_args(family)
    for i in range(first_page - 1, total):
        page = doc[i]
        r = page.rect
        y = margin + size if vpos == "top" else r.height - margin
        text = fmt.format(n=i - first_page + 1 + start, total=total, date=today)
        box = pymupdf.Rect(margin, y - size - 2, r.width - margin, y + 4)
        page.insert_textbox(to_unrotated(page, box), text, fontsize=size, align=align,
                            color=color, rotate=page.rotation, **kw)


def header_footer(doc, header="", footer="", size=10, family="DejaVu Sans", color=(0, 0, 0)):
    if header:
        page_numbers(doc, header, "top-center", size=size, family=family, color=color)
    if footer:
        page_numbers(doc, footer, "bottom-center", size=size, family=family, color=color)


def redact(page, rect, fill=(0, 0, 0)):
    page.add_redact_annot(to_unrotated(page, rect), fill=fill)
    page.apply_redactions()


def redact_text(doc, needle, fill=(0, 0, 0)):
    count = 0
    for page in doc:
        hits = page.search_for(needle)
        for r in hits:
            page.add_redact_annot(r, fill=fill)
        if hits:
            page.apply_redactions()
            count += len(hits)
    return count


@dataclass
class IdHit:
    page: int
    value: str
    cells: list
    source: str = "text"

    @property
    def rect(self):
        r = pymupdf.Rect(self.cells[0])
        for c in self.cells[1:]:
            r |= c
        return r


def _ocr_variants(img):
    from PIL import ImageOps
    gray = ImageOps.grayscale(img)
    blue = ImageOps.autocontrast(img.convert("RGB").split()[2])
    return [gray.point(lambda x: 255 if x > 120 else 0),
            blue.point(lambda x: 255 if x > 110 else 0),
            ImageOps.autocontrast(gray)]


def _tesseract_words(img, language):
    with tempfile.TemporaryDirectory(prefix="acrpdf-ocr-") as tmp:
        path = os.path.join(tmp, "page.png")
        img.save(path)
        out = subprocess.run(["tesseract", path, "stdout", "-l", language, "--psm", "11", "tsv"],
                             capture_output=True, text=True, timeout=300).stdout
    words = []
    for line in out.splitlines()[1:]:
        cols = line.split("\t")
        if len(cols) == 12 and cols[11].strip():
            words.append((cols[11].strip(), int(cols[6]), int(cols[7]), int(cols[8]), int(cols[9])))
    return words


def _ocr_id_hits(page, language):
    from PIL import Image
    regions = []
    for info in page.get_image_info():
        r = pymupdf.Rect(info["bbox"]) & page.rect
        if r.width > 20 and r.height > 20:
            native = info["width"] / (pymupdf.Rect(info["bbox"]).width / 72)
            regions.append((r, int(min(500, max(200, native)))))
    regions = regions or [(page.rect, 300)]
    hits = []
    for clip, dpi in regions:
        pix = page.get_pixmap(clip=clip, dpi=dpi)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        scale = 72 / dpi
        for variant in _ocr_variants(img):
            for text, left, top, width, height in _tesseract_words(variant, language):
                for start, end in dni.find_in_token(text):
                    cell_w = width / len(text)
                    cells = [pymupdf.Rect(clip.x0 + (left + k * cell_w) * scale, clip.y0 + top * scale,
                                          clip.x0 + (left + (k + 1) * cell_w) * scale, clip.y0 + (top + height) * scale)
                             * page.derotation_matrix for k in range(start, end)]
                    value = text[start:end].upper()
                    hit = IdHit(page.number, value, cells, "ocr")
                    if not any(h.value == value and h.rect.intersects(hit.rect) for h in hits):
                        hits.append(hit)
    return hits


def scan_ids(doc, validate=True, ocr="auto", language="spa"):
    use_ocr = ocr and shutil.which("tesseract") is not None
    hits = []
    for page in doc:
        for value in dni.find(page.get_text(), validate):
            for r in page.search_for(value):
                hits.append(IdHit(page.number, value, [r], "text"))
        if use_ocr and (ocr == "always" or page.get_image_info()):
            for hit in _ocr_id_hits(page, language):
                if not any(h.page == hit.page and h.rect.intersects(hit.rect) for h in hits):
                    hits.append(hit)
    return hits


def apply_id_hits(doc, hits, mode="mask"):
    by_page = {}
    for hit in hits:
        by_page.setdefault(hit.page, []).append(hit)
    for number, page_hits in by_page.items():
        page = doc[number]
        rewrite = []
        sizes = {tuple(pymupdf.Rect(sp["bbox"])): sp["size"] for b in page.get_text("dict")["blocks"]
                 for line in b.get("lines", []) for sp in line["spans"]}
        for hit in page_hits:
            if hit.source == "ocr":
                keep = set(dni.visible_positions(hit.value)) if mode == "mask" else set()
                offset = len(hit.value) - len(dni.normalize(hit.value))
                for k, cell in enumerate(hit.cells):
                    if k - offset not in keep:
                        page.add_redact_annot(cell, fill=(0, 0, 0))
            else:
                page.add_redact_annot(hit.rect, fill=(0, 0, 0) if mode == "black" else None)
                if mode == "mask":
                    size = next((sz for bbox, sz in sizes.items() if pymupdf.Rect(bbox).intersects(hit.rect)),
                                max(5.0, hit.rect.height * 0.78))
                    rewrite.append((hit.rect, dni.mask(hit.value), size))
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
        for rect, masked, size in rewrite:
            page.insert_text((rect.x0, rect.y1 - rect.height * 0.22), masked, fontsize=size,
                             fontname="helv", color=(0, 0, 0))
    return len(hits)


@dataclass
class IdWatermark:
    text: str = "SOLO PARA {purpose} · {date}"
    purpose: str = ""
    size: float = 16
    stroke: float = 0.0
    opacity: float = 0.35
    color: tuple = (0.12, 0.25, 0.55)
    angle: float = 30
    spacing: float = 1.0
    footer: str = "{doc} para {purpose}"
    footer_size: float = 11
    document: str = "DNI"
    only_images: bool = False

    def render(self, template):
        date = datetime.date.today().strftime("%d/%m/%Y")
        try:
            return template.format(purpose=self.purpose or "…", date=date, doc=self.document).strip()
        except (KeyError, IndexError, ValueError):
            return template.strip()


def _watermark_pattern(width, height, wm, center):
    text = wm.render(wm.text)
    kw = font_args("DejaVu Sans", bold=True)
    font = font_object("DejaVu Sans", bold=True)
    unit = text + "      "
    unit_w = max(font.text_length(unit, fontsize=wm.size), 20)
    reach = width + height
    line = unit * (int(reach * 2 / unit_w) + 2)
    step = wm.size * 3.2 * wm.spacing
    doc = pymupdf.open()
    page = doc.new_page(width=width, height=height)
    shape = page.new_shape()
    y, row = center.y - reach, 0
    while y < center.y + reach:
        x = center.x - reach - (row % 2) * unit_w / 2
        shape.insert_text(pymupdf.Point(x, y), line, fontsize=wm.size, color=wm.color, fill=wm.color,
                          fill_opacity=wm.opacity, stroke_opacity=wm.opacity,
                          render_mode=2 if wm.stroke > 0 else 0, border_width=wm.stroke / max(wm.size, 1),
                          morph=(center, pymupdf.Matrix(-wm.angle)), **kw)
        y += step
        row += 1
    shape.commit()
    return doc


def apply_id_watermark(page, wm):
    r = page.rect
    areas = [pymupdf.Rect(i["bbox"]) * page.rotation_matrix & r for i in page.get_image_info()] \
        if wm.only_images else []
    areas = [a for a in areas if a.width > 30 and a.height > 30] or [r]
    for area in areas:
        center = pymupdf.Point((area.x0 + area.x1) / 2, (area.y0 + area.y1) / 2)
        pattern = _watermark_pattern(r.width, r.height, wm, center)
        page.show_pdf_page(to_unrotated(page, area), pattern, 0, clip=area, rotate=page.rotation, overlay=True)
        pattern.close()
    if wm.footer:
        kw = {"fontname": "hebo"}
        h = wm.footer_size * 2.2
        band = pymupdf.Rect(0, r.height - h, r.width, r.height)
        page.draw_rect(to_unrotated(page, band), color=None, fill=(1, 1, 1), fill_opacity=0.85, overlay=True)
        page.insert_textbox(to_unrotated(page, band + (10, h * 0.22, -10, 0)), wm.render(wm.footer),
                            fontsize=wm.footer_size, align=1, color=wm.color, rotate=page.rotation, **kw)


def flatten_to_images(doc, dpi=200, quality=85):
    out = pymupdf.open()
    for page in doc:
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, stream=pix.tobytes("jpeg", jpg_quality=quality))
    return out


def protect_identity(data, hits=(), mode="mask", watermark=None, flatten=True, dpi=200, pages=None):
    doc = pymupdf.open("pdf", data)
    if pages is not None:
        doc.select(list(pages))
        index = {old: new for new, old in enumerate(pages)}
        hits = [IdHit(index[h.page], h.value, h.cells, h.source) for h in hits if h.page in index]
    apply_id_hits(doc, hits, mode)
    if watermark is not None:
        for page in doc:
            apply_id_watermark(page, watermark)
    if flatten:
        flat = flatten_to_images(doc, dpi)
        doc.close()
        doc = flat
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out


def find_ids(doc, validate=True, ocr="auto"):
    return [(h.page, h.value) for h in scan_ids(doc, validate, ocr)]


def hide_ids(doc, mode="mask", validate=True, only=None, ocr="auto"):
    hits = [h for h in scan_ids(doc, validate, ocr) if only is None or h.value in only]
    return apply_id_hits(doc, hits, mode)


FIELD_TYPES = {
    "text": pymupdf.PDF_WIDGET_TYPE_TEXT,
    "check": pymupdf.PDF_WIDGET_TYPE_CHECKBOX,
    "radio": pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON,
    "combo": pymupdf.PDF_WIDGET_TYPE_COMBOBOX,
    "list": pymupdf.PDF_WIDGET_TYPE_LISTBOX,
    "signature": pymupdf.PDF_WIDGET_TYPE_SIGNATURE,
}


def unique_field_name(doc, base):
    names = {w.field_name for p in doc for w in p.widgets()}
    i = 1
    while f"{base}{i}" in names:
        i += 1
    return f"{base}{i}"


def add_field(page, kind, rect, name, options=None, fontsize=0, multiline=False, required=False,
              tooltip=""):
    w = pymupdf.Widget()
    w.field_type = FIELD_TYPES[kind]
    w.field_name = name
    w.rect = to_unrotated(page, rect)
    w.field_label = tooltip or name
    if kind in ("text", "combo", "list"):
        w.text_font = "Helv"
        w.text_fontsize = fontsize
        w.border_color = (0.45, 0.5, 0.65)
        w.fill_color = (0.91, 0.94, 1)
    if kind in ("combo", "list"):
        w.choice_values = options or ["1", "2"]
        w.field_value = w.choice_values[0]
    if kind in ("check", "radio"):
        w.field_value = False
        w.border_color = (0.3, 0.3, 0.3)
    flags = 0
    if multiline and kind == "text":
        flags |= pymupdf.PDF_TX_FIELD_IS_MULTILINE
    if required:
        flags |= pymupdf.PDF_FIELD_IS_REQUIRED
    w.field_flags = flags
    page.add_widget(w)
    return w


def find_widget(page, xref):
    for w in page.widgets():
        if w.xref == xref:
            return w
    return None


def set_widget_rect(page, xref, rect):
    w = find_widget(page, xref)
    if w is not None:
        w.rect = to_unrotated(page, rect)
        w.update()


def set_widget_fontsize(page, xref, size):
    w = find_widget(page, xref)
    if w is not None:
        w.text_fontsize = size
        w.update()


def all_fields(doc):
    return [(p.number, w.field_name, w.field_type_string, w.field_value, w.xref)
            for p in doc for w in p.widgets()]


def set_field(doc, page_no, xref, value):
    page = doc[page_no]
    w = find_widget(page, xref)
    if w is None:
        return False
    if w.field_type in (pymupdf.PDF_WIDGET_TYPE_CHECKBOX, pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON):
        w.field_value = bool(value)
    else:
        w.field_value = value
    w.update()
    return True


def delete_widget(doc, page_no, xref):
    page = doc[page_no]
    w = find_widget(page, xref)
    if w is not None:
        page.delete_widget(w)


def flatten(doc):
    doc.bake(annots=True, widgets=True)


def has_signatures(doc):
    return any(w.is_signed for p in doc for w in p.widgets(types=[pymupdf.PDF_WIDGET_TYPE_SIGNATURE]))


def rotate_pages(doc, pages, angle):
    for i in pages:
        doc[i].set_rotation((doc[i].rotation + angle) % 360)


def delete_pages(doc, pages):
    doc.delete_pages(sorted(set(pages)))


def duplicate_pages(doc, pages):
    for offset, i in enumerate(sorted(pages)):
        doc.fullcopy_page(i + offset, i + offset + 1)


def insert_blank(doc, after):
    ref = doc[min(after, doc.page_count - 1)].rect if doc.page_count else pymupdf.paper_rect("a4")
    doc.new_page(after + 1, width=ref.width, height=ref.height)


def reorder(doc, order):
    doc.select(order)


def insert_files(doc, paths, at=None):
    for path in paths:
        src = open_any(path)
        doc.insert_pdf(src, start_at=-1 if at is None else at)
        if at is not None:
            at += src.page_count
        src.close()


def extract_pages(doc, pages, out_path):
    new = pymupdf.open()
    for i in pages:
        new.insert_pdf(doc, from_page=i, to_page=i)
    new.save(out_path, garbage=3, deflate=True)
    new.close()


def crop_page(page, rect):
    r = to_unrotated(page, rect)
    page.set_cropbox(r + (page.cropbox.x0, page.cropbox.y0, page.cropbox.x0, page.cropbox.y0))


def crop_pages(doc, rect, pages=None):
    for i in pages if pages is not None else range(doc.page_count):
        page = doc[i]
        crop_page(page, pymupdf.Rect(rect) & page.rect)


def parse_ranges(spec, page_count):
    groups = []
    for part in spec.replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            a = int(a) if a.strip() else 1
            b = int(b) if b.strip() else page_count
        else:
            a = b = int(part)
        if not (1 <= a <= b <= page_count):
            raise OperationError(_("Invalid page range: %s") % part)
        groups.append(list(range(a - 1, b)))
    return groups


def _soffice():
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise OperationError(_("LibreOffice is not installed (needed to convert office documents)."))
    return exe


def office_to_pdf(path, out_path=None):
    out_dir = tempfile.mkdtemp(prefix="acrpdf-")
    subprocess.run([_soffice(), "--headless", "--norestore", "--convert-to", "pdf", "--outdir", out_dir, path],
                   check=True, capture_output=True, timeout=600)
    pdf = os.path.join(out_dir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    if not os.path.exists(pdf):
        raise OperationError(_("LibreOffice could not convert %s") % os.path.basename(path))
    if out_path:
        shutil.move(pdf, out_path)
        return out_path
    return pdf


def open_any(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return pymupdf.open(path)
    if ext in IMAGE_EXT or ext in EBOOK_EXT:
        src = pymupdf.open(path)
        data = src.convert_to_pdf()
        src.close()
        return pymupdf.open("pdf", data)
    return pymupdf.open(office_to_pdf(path))


def create_from_files(paths):
    doc = pymupdf.open()
    for p in paths:
        src = open_any(p)
        doc.insert_pdf(src)
        src.close()
    return doc


SAVE_COMPRESSION = {
    "lossless": None,
    "high": (330, 300, 90),
    "recommended": (230, 200, 80),
    "small": (170, 150, 65),
}


def compact_bytes(data, level="recommended"):
    doc = pymupdf.open("pdf", data)
    try:
        if doc.is_encrypted or doc.needs_pass or has_signatures(doc):
            return data
        images = SAVE_COMPRESSION.get(level)
        if images:
            threshold, target, quality = images
            try:
                doc.rewrite_images(dpi_threshold=threshold, dpi_target=target, quality=quality)
            except Exception:
                pass
        try:
            doc.subset_fonts()
        except Exception:
            pass
        out = doc.tobytes(garbage=3, deflate=True, deflate_images=True, deflate_fonts=True, use_objstms=1)
    finally:
        doc.close()
    return out if len(out) < len(data) else data


def save_doc(doc, out_path):
    doc.save(out_path, garbage=3, deflate=True)
    return out_path


def merge(paths, out_path):
    doc = create_from_files(paths)
    save_doc(doc, out_path)
    doc.close()
    return [out_path]


def split(path, out_dir, ranges=None, every=None):
    doc = pymupdf.open(path)
    stem = os.path.splitext(os.path.basename(path))[0]
    if ranges:
        groups = parse_ranges(ranges, doc.page_count)
    else:
        n = max(1, every or 1)
        groups = [list(range(i, min(i + n, doc.page_count))) for i in range(0, doc.page_count, n)]
    outputs = []
    for g in groups:
        name = f"{stem}_{g[0] + 1}" + (f"-{g[-1] + 1}" if len(g) > 1 else "") + ".pdf"
        out = os.path.join(out_dir, name)
        extract_pages(doc, g, out)
        outputs.append(out)
    doc.close()
    return outputs


def remove_pages(path, out_path, spec):
    doc = pymupdf.open(path)
    drop = sorted({i for g in parse_ranges(spec, doc.page_count) for i in g})
    if len(drop) >= doc.page_count:
        raise OperationError(_("Cannot remove every page."))
    doc.delete_pages(drop)
    save_doc(doc, out_path)
    return [out_path]


def extract(path, out_path, spec):
    doc = pymupdf.open(path)
    pages = [i for g in parse_ranges(spec, doc.page_count) for i in g]
    extract_pages(doc, pages, out_path)
    return [out_path]


COMPRESSION_LEVELS = {"extreme": ("/screen", 72), "recommended": ("/ebook", 150), "low": ("/printer", 300)}


def compress(path, out_path, level="recommended"):
    preset, dpi = COMPRESSION_LEVELS[level]
    candidates = []
    tmp = tempfile.mkdtemp(prefix="acrpdf-")
    gs = shutil.which("gs")
    if gs:
        gs_out = os.path.join(tmp, "gs.pdf")
        r = subprocess.run([gs, "-q", "-dSAFER", "-dBATCH", "-dNOPAUSE", "-sDEVICE=pdfwrite", "-dAutoRotatePages=/None",
                            "-dCompatibilityLevel=1.7", f"-dPDFSETTINGS={preset}",
                            "-dDetectDuplicateImages=true", f"-dColorImageResolution={dpi}",
                            f"-dGrayImageResolution={dpi}", "-dDownsampleColorImages=true",
                            "-dDownsampleGrayImages=true", f"-sOutputFile={gs_out}", path],
                           capture_output=True, timeout=900)
        if r.returncode == 0 and os.path.exists(gs_out) and _opens(gs_out):
            candidates.append(gs_out)
    mu_out = os.path.join(tmp, "mu.pdf")
    doc = pymupdf.open(path)
    try:
        doc.rewrite_images(dpi_threshold=dpi + 10, dpi_target=dpi, quality={72: 50, 150: 70, 300: 85}[dpi])
    except Exception:
        pass
    doc.save(mu_out, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True,
             use_objstms=1)
    doc.close()
    candidates.append(mu_out)
    best = min(candidates + [path], key=os.path.getsize)
    shutil.copyfile(best, out_path)
    shutil.rmtree(tmp, ignore_errors=True)
    return [out_path]


def _opens(path):
    try:
        with pymupdf.open(path) as d:
            return d.page_count > 0
    except Exception:
        return False


def repair(path, out_path):
    try:
        doc = pymupdf.open(path)
        doc.save(out_path, garbage=4, deflate=True, clean=True)
        doc.close()
        if _opens(out_path):
            return [out_path]
    except Exception:
        pass
    gs = shutil.which("gs")
    if not gs:
        raise OperationError(_("The file could not be repaired."))
    subprocess.run([gs, "-q", "-dSAFER", "-dBATCH", "-dNOPAUSE", "-sDEVICE=pdfwrite", "-dAutoRotatePages=/None",
                    f"-sOutputFile={out_path}", path], check=True, capture_output=True, timeout=900)
    return [out_path]


def ocr_languages():
    exe = shutil.which("tesseract")
    if not exe:
        return []
    out = subprocess.run([exe, "--list-langs"], capture_output=True, text=True).stdout
    return [line.strip() for line in out.splitlines()[1:] if line.strip() and line.strip() != "osd"]


def _tessdata():
    for base in glob.glob("/usr/share/tesseract-ocr/*/tessdata") + ["/usr/share/tessdata"]:
        if os.path.isdir(base):
            return base
    return None


def ocr(path, out_path, language="spa+eng", dpi=300):
    if not ocr_languages():
        raise OperationError(_("Tesseract OCR is not installed. Install the packages "
                               "tesseract-ocr and tesseract-ocr-spa."))
    tessdata = _tessdata()
    src = pymupdf.open(path)
    out = pymupdf.open()
    for page in src:
        if len(page.get_text().strip()) > 20:
            out.insert_pdf(src, from_page=page.number, to_page=page.number)
            continue
        pix = page.get_pixmap(dpi=dpi)
        ocr_pdf = pymupdf.open("pdf", pix.pdfocr_tobytes(language=language, tessdata=tessdata))
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.show_pdf_page(new.rect, ocr_pdf, 0)
        ocr_pdf.close()
    out.save(out_path, garbage=3, deflate=True)
    return [out_path]


def images_to_pdf(paths, out_path, page_size="a4", margin=20, landscape=False):
    doc = pymupdf.open()
    for p in paths:
        with open(p, "rb") as f:
            data = f.read()
        iw, ih = graphic_size(data)
        if page_size == "fit":
            page = doc.new_page(width=iw * 0.75 + 2 * margin, height=ih * 0.75 + 2 * margin)
        else:
            r = pymupdf.paper_rect(page_size + ("-l" if landscape else ""))
            page = doc.new_page(width=r.width, height=r.height)
        box = page.rect + (margin, margin, -margin, -margin)
        insert_graphic(page, box, data)
    save_doc(doc, out_path)
    return [out_path]


def pdf_to_images(path, out_dir, fmt="jpg", dpi=150, mode="pages"):
    doc = pymupdf.open(path)
    stem = os.path.splitext(os.path.basename(path))[0]
    outputs = []
    if mode == "extract":
        seen = set()
        for page in doc:
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen:
                    continue
                seen.add(xref)
                info = doc.extract_image(xref)
                out = os.path.join(out_dir, f"{stem}_img{len(seen):03d}.{info['ext']}")
                with open(out, "wb") as f:
                    f.write(info["image"])
                outputs.append(out)
        return outputs
    for page in doc:
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        out = os.path.join(out_dir, f"{stem}_{page.number + 1:03d}.{fmt}")
        if fmt in ("jpg", "jpeg"):
            pix.save(out, jpg_quality=90)
        else:
            pix.save(out)
        outputs.append(out)
    return outputs


def pdf_to_word(path, out_path):
    try:
        import logging
        from pdf2docx import Converter
        logging.getLogger("root").setLevel(logging.ERROR)
    except ImportError:
        return [_libreoffice_export(path, out_path, "docx", "writer_pdf_import")]
    conv = Converter(path)
    try:
        conv.convert(out_path)
    finally:
        conv.close()
    return [out_path]


def _libreoffice_export(path, out_path, ext, infilter):
    tmp = tempfile.mkdtemp(prefix="acrpdf-")
    subprocess.run([_soffice(), "--headless", "--norestore", f"--infilter={infilter}", "--convert-to", ext,
                    "--outdir", tmp, path], check=True, capture_output=True, timeout=600)
    produced = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + "." + ext)
    shutil.move(produced, out_path)
    return out_path


def pdf_to_excel(path, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    doc = pymupdf.open(path)
    for page in doc:
        ws = wb.create_sheet(_("Page %d") % (page.number + 1))
        rows = []
        try:
            tables = page.find_tables().tables
        except Exception:
            tables = []
        for t in tables:
            rows.extend(t.extract())
            rows.append([])
        if not rows:
            rows = _text_rows(page)
        for row in rows:
            ws.append([("" if c is None else str(c)) for c in row])
        if rows and tables:
            for cell in ws[1]:
                cell.font = Font(bold=True)
    if not wb.sheetnames:
        wb.create_sheet("1")
    wb.save(out_path)
    return [out_path]


def _text_rows(page, gap=14):
    lines = {}
    for w in page.get_text("words"):
        lines.setdefault((w[5], w[6]), []).append(w)
    rows = []
    for key in sorted(lines, key=lambda k: (lines[k][0][1], lines[k][0][0])):
        words = sorted(lines[key], key=lambda w: w[0])
        cells, current, last_x = [], [], None
        for w in words:
            if last_x is not None and w[0] - last_x > gap:
                cells.append(" ".join(current))
                current = []
            current.append(w[4])
            last_x = w[2]
        cells.append(" ".join(current))
        rows.append(cells)
    return rows


def pdf_to_powerpoint(path, out_path, dpi=150):
    from pptx import Presentation
    from pptx.util import Emu
    doc = pymupdf.open(path)
    prs = Presentation()
    first = doc[0].rect
    prs.slide_width = Emu(int(first.width * 12700))
    prs.slide_height = Emu(int(first.height * 12700))
    blank = prs.slide_layouts[6]
    for page in doc:
        slide = prs.slides.add_slide(blank)
        png = page.get_pixmap(dpi=dpi).tobytes("png")
        slide.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
        slide.notes_slide.notes_text_frame.text = page.get_text()
    prs.save(out_path)
    return [out_path]


def pdf_to_text(path, out_path):
    doc = pymupdf.open(path)
    with open(out_path, "w", encoding="utf-8") as f:
        for page in doc:
            f.write(page.get_text())
            f.write("\f")
    return [out_path]


def _pdfa_definition(tmp):
    gs_lib = sorted(glob.glob("/usr/share/ghostscript/*/lib/PDFA_def.ps"))
    icc = next((p for p in ("/usr/share/color/icc/ghostscript/srgb.icc",
                            "/usr/share/color/icc/colord/sRGB.icc") if os.path.exists(p)), None)
    if not gs_lib or not icc:
        raise OperationError(_("Ghostscript PDF/A resources were not found."))
    with open(gs_lib[-1], encoding="latin-1") as f:
        text = f.read()
    text = re.sub(r"^/ICCProfile \(.*?\)", f"/ICCProfile ({icc})", text, flags=re.M)
    out = os.path.join(tmp, "PDFA_def.ps")
    with open(out, "w", encoding="latin-1") as f:
        f.write(text)
    return out, os.path.dirname(icc)


def pdf_to_pdfa(path, out_path, level=2):
    gs = shutil.which("gs")
    if not gs:
        raise OperationError(_("Ghostscript is not installed."))
    tmp = tempfile.mkdtemp(prefix="acrpdf-")
    definition, icc_dir = _pdfa_definition(tmp)
    subprocess.run([gs, "-q", f"-dPDFA={level}", "-dBATCH", "-dNOPAUSE", "-dNOOUTERSAVE",
                    "-dPDFACompatibilityPolicy=1", "-sColorConversionStrategy=RGB",
                    "-sDEVICE=pdfwrite", "-dAutoRotatePages=/None", f"--permit-file-read={icc_dir}/",
                    f"-sOutputFile={out_path}", definition, path],
                   check=True, capture_output=True, timeout=900)
    shutil.rmtree(tmp, ignore_errors=True)
    return [out_path]


def rotate_file(path, out_path, angle=90, spec=""):
    doc = pymupdf.open(path)
    pages = [i for g in parse_ranges(spec, doc.page_count) for i in g] if spec.strip() else range(doc.page_count)
    rotate_pages(doc, pages, angle)
    save_doc(doc, out_path)
    return [out_path]


def protect(path, out_path, user_pw, owner_pw=None, allow_print=True, allow_copy=False,
            allow_edit=False, password=None):
    doc = pymupdf.open(path)
    if doc.needs_pass and not doc.authenticate(password or ""):
        raise OperationError(_("Wrong password."))
    perm = pymupdf.PDF_PERM_ACCESSIBILITY
    if allow_print:
        perm |= pymupdf.PDF_PERM_PRINT | pymupdf.PDF_PERM_PRINT_HQ
    if allow_copy:
        perm |= pymupdf.PDF_PERM_COPY
    if allow_edit:
        perm |= pymupdf.PDF_PERM_MODIFY | pymupdf.PDF_PERM_ANNOTATE | pymupdf.PDF_PERM_FORM
    doc.save(out_path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=user_pw or "",
             owner_pw=owner_pw or user_pw, permissions=perm, garbage=3, deflate=True)
    return [out_path]


def unlock(path, out_path, password=""):
    doc = pymupdf.open(path)
    if doc.needs_pass and not doc.authenticate(password):
        raise OperationError(_("Wrong password."))
    doc.save(out_path, encryption=pymupdf.PDF_ENCRYPT_NONE, garbage=3, deflate=True)
    return [out_path]


def compare(path_a, path_b, out_path):
    a, b = pymupdf.open(path_a), pymupdf.open(path_b)
    changes = 0
    for i, page in enumerate(b):
        words_b = page.get_text("words")
        words_a = a[i].get_text("words") if i < a.page_count else []
        sm = difflib.SequenceMatcher(None, [w[4] for w in words_a], [w[4] for w in words_b], autojunk=False)
        removed = []
        for op, a0, a1, b0, b1 in sm.get_opcodes():
            if op in ("replace", "insert"):
                for w in words_b[b0:b1]:
                    h = page.add_highlight_annot(pymupdf.Rect(w[:4]))
                    h.set_colors(stroke=(0.55, 0.9, 0.55))
                    h.update()
                changes += b1 - b0
            if op in ("replace", "delete"):
                removed.append(" ".join(w[4] for w in words_a[a0:a1]))
        if removed:
            note = page.add_text_annot((page.rect.width - 30, 20), _("Removed:") + "\n" + "\n".join(removed),
                                       icon="Note")
            note.set_colors(stroke=(0.9, 0.2, 0.2))
            note.update()
            changes += len(removed)
    b.save(out_path, garbage=3, deflate=True)
    return [out_path], changes


def numbers_file(path, out_path, **kwargs):
    doc = pymupdf.open(path)
    page_numbers(doc, **kwargs)
    save_doc(doc, out_path)
    return [out_path]


def watermark_file(path, out_path, **kwargs):
    doc = pymupdf.open(path)
    watermark(doc, **kwargs)
    save_doc(doc, out_path)
    return [out_path]


def crop_margins_file(path, out_path, margin=0):
    doc = pymupdf.open(path)
    for page in doc:
        boxes = [pymupdf.Rect(b[:4]) for b in page.get_text("blocks")]
        boxes += [d["rect"] for d in page.get_drawings()]
        boxes += [page.get_image_bbox(img) for img in page.get_images(full=True)]
        boxes = [r for r in boxes if r.is_valid and not r.is_empty and not r.is_infinite]
        if not boxes:
            continue
        r = boxes[0]
        for x in boxes[1:]:
            r |= x
        r = (r + (-margin, -margin, margin, margin)) & page.rect
        page.set_cropbox(r + (page.cropbox.x0, page.cropbox.y0, page.cropbox.x0, page.cropbox.y0))
    save_doc(doc, out_path)
    return [out_path]


def excel_to_pdf(path, out_path, landscape=False, fit_width=True):
    from openpyxl import load_workbook
    tmp = tempfile.mkdtemp(prefix="acrpdf-")
    src = path
    if not path.lower().endswith((".xlsx", ".xlsm")):
        subprocess.run([_soffice(), "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", tmp, path],
                       check=True, capture_output=True, timeout=600)
        src = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
    wb = load_workbook(src)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        if fit_width:
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
    prepared = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
    wb.save(prepared)
    office_to_pdf(prepared, out_path)
    shutil.rmtree(tmp, ignore_errors=True)
    return [out_path]


def _chrome():
    for name in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "microsoft-edge"):
        exe = shutil.which(name)
        if exe:
            return exe
    return None


def html_to_pdf(source, out_path):
    is_url = re.match(r"^https?://", source, re.I) is not None
    chrome = _chrome()
    if chrome:
        profile = tempfile.mkdtemp(prefix="acrpdf-chrome-")
        target = source if is_url else "file://" + os.path.abspath(source)
        args = [chrome, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                f"--user-data-dir={profile}", f"--print-to-pdf={out_path}", target]
        r = subprocess.run(args, capture_output=True, timeout=180)
        shutil.rmtree(profile, ignore_errors=True)
        if r.returncode == 0 and os.path.exists(out_path):
            return [out_path]
    if is_url:
        raise OperationError(_("Converting web addresses needs Google Chrome or Chromium installed."))
    return [office_to_pdf(source, out_path)]


def pdf_to_markdown(path, out_path):
    import contextlib
    import pymupdf4llm
    with contextlib.redirect_stdout(io.StringIO()):
        text = pymupdf4llm.to_markdown(path)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(text)
    return [out_path]


def document_text(path_or_doc, max_chars=None):
    doc = pymupdf.open(path_or_doc) if isinstance(path_or_doc, str) else path_or_doc
    parts = []
    for page in doc:
        parts.append(f"[{page.number + 1}]\n" + page.get_text())
        if max_chars and sum(len(p) for p in parts) > max_chars:
            break
    return "\n".join(parts)


def text_blocks(page):
    blocks = []
    for b in page.get_text("dict", flags=pymupdf.TEXTFLAGS_TEXT)["blocks"]:
        if b.get("type") != 0:
            continue
        lines, sizes, color = [], [], 0
        for line in b["lines"]:
            lines.append("".join(s["text"] for s in line["spans"]))
            for s in line["spans"]:
                sizes.append(s["size"])
                color = s["color"]
        text = "\n".join(lines).strip()
        if text:
            blocks.append((pymupdf.Rect(b["bbox"]), text, sorted(sizes)[len(sizes) // 2], color))
    return blocks


def translate_pdf(path, out_path, translate, progress=None):
    doc = pymupdf.open(path)
    pages = [(page.number, text_blocks(page)) for page in doc]
    texts = [t for _n, blocks in pages for _r, t, _s, _c in blocks]
    translated = iter(translate(texts)) if texts else iter(())
    for number, blocks in pages:
        if not blocks:
            continue
        page = doc[number]
        for rect, _t, _s, _c in blocks:
            page.add_redact_annot(rect, fill=None)
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
        for rect, _t, size, color in blocks:
            new = next(translated)
            html = new.replace("&", "&amp;").replace("<", "&lt;").replace("\n", "<br>")
            css = f"* {{font-family: sans-serif; font-size: {size:.1f}px; color: #{color:06x};}}"
            page.insert_htmlbox(rect + (0, 0, 2, 2), html, css=css)
        if progress:
            progress(number + 1, doc.page_count)
    doc.save(out_path, garbage=3, deflate=True)
    return [out_path]


def detect_fields(doc, min_line=40):
    found = []
    for page in doc:
        rot = page.rotation_matrix
        words = page.get_text("words")
        occupied = [pymupdf.Rect(w[:4]) for w in words if not set(w[4]) <= set("_.…")]
        for w in words:
            r = pymupdf.Rect(w[:4])
            if set(w[4]) <= set("_") and r.width >= min_line:
                found.append((page.number, "text", pymupdf.Rect(r.x0, r.y0 - 4, r.x1, r.y1) * rot))
            elif w[4] in ("☐", "□", "❑", "❒"):
                found.append((page.number, "check", r * rot))
        for d in page.get_drawings():
            for item in d["items"]:
                if item[0] == "l":
                    p1, p2 = item[1], item[2]
                    if abs(p1.y - p2.y) < 1 and abs(p2.x - p1.x) >= min_line:
                        line = pymupdf.Rect(min(p1.x, p2.x), p1.y - 16, max(p1.x, p2.x), p1.y - 1)
                        if not any(line.intersects(o) for o in occupied):
                            found.append((page.number, "text", line * rot))
                elif item[0] == "re":
                    r = pymupdf.Rect(item[1])
                    if 6 <= r.width <= 20 and 6 <= r.height <= 20 and abs(r.width - r.height) < 3:
                        found.append((page.number, "check", r * rot))
                    elif r.width >= min_line and 12 <= r.height <= 40 and \
                            not any(r.intersects(o) for o in occupied):
                        found.append((page.number, "text", r * rot))
    unique = []
    for item in found:
        if not any(u[0] == item[0] and (u[2] & item[2]).get_area() > 0.5 * item[2].get_area() for u in unique):
            unique.append(item)
    existing = {(p.number, round(w.rect.x0), round(w.rect.y0)) for p in doc for w in p.widgets()}
    return [u for u in unique if (u[0], round(u[2].x0), round(u[2].y0)) not in existing]


def add_detected_fields(doc, detections):
    for page_number, kind, rect in detections:
        base = "Text" if kind == "text" else "Check"
        add_field(doc[page_number], kind, rect, unique_field_name(doc, base))
    return len(detections)

